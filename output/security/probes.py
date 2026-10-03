"""Local audit probes. Mocks replace the database, email and external services.

These assertions document the audited behavior, including vulnerabilities;
they are evidence for this review, not regression tests expecting secure behavior.
Run from the repository root: .codex-venv/Scripts/python.exe output/security/probes.py
"""
import io
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
os.environ["INIT_DB"] = "false"
os.environ["SECRET_KEY"] = "isolated-security-audit-key-not-for-deployment"

import app
from flask import session
from openpyxl import Workbook, load_workbook
from repositories import notificaciones
from werkzeug.datastructures import FileStorage


class SecurityReviewEvidence(unittest.TestCase):
    def client(self):
        return app.app.test_client()

    def admin(self, client):
        with client.session_transaction() as state:
            state.update(user_id=999, username="deleted-admin", rol="admin", permisos=[],
                         _csrf_token="audit-token", _permanent=True)

    def test_revoked_admin_session_still_reads_users(self):
        client = self.client()
        self.admin(client)
        conn = MagicMock()
        conn.execute.return_value.fetchall.return_value = []
        with patch.object(app, "get_connection", return_value=conn), patch.object(
            app, "obtener_config_mantenimiento", return_value={"activo": False}
        ):
            response = client.get("/usuarios")
        self.assertEqual(response.status_code, 200)
        queries = [call.args[0] for call in conn.execute.call_args_list]
        self.assertEqual(len(queries), 1)
        self.assertIn("ORDER BY username", queries[0])

    def test_session_cookie_can_be_replayed_after_logout(self):
        client = self.client()
        self.admin(client)
        old_cookie = client.get_cookie(app.app.config["SESSION_COOKIE_NAME"]).value
        with patch.object(app, "registrar_auditoria"):
            self.assertEqual(client.get("/logout").status_code, 302)
        client.set_cookie(app.app.config["SESSION_COOKIE_NAME"], old_cookie)
        conn = MagicMock()
        conn.execute.return_value.fetchall.return_value = []
        with patch.object(app, "get_connection", return_value=conn), patch.object(
            app, "obtener_config_mantenimiento", return_value={"activo": False}
        ):
            self.assertEqual(client.get("/usuarios").status_code, 200)

    def test_explicit_empty_permissions_reenable_role_defaults(self):
        with app.app.test_request_context("/"):
            session.update(user_id=1, rol="entrenador", permisos=[])
            self.assertTrue(app.tiene_permiso("jugadores_gestionar"))

    def test_dni_alone_discloses_portal_bearer_token(self):
        client = self.client()
        with client.session_transaction() as state:
            state["_csrf_token"] = "audit-token"
        conn = MagicMock()
        conn.execute.return_value.fetchall.return_value = [{"id": 123, "portal_token": "fake-player-bearer"}]
        with patch.object(app, "get_connection", return_value=conn), patch.object(
            app, "consumir_limite_publico"
        ) as limiter:
            response = client.post("/portal", data={"identificador": "12345678", "_csrf_token": "audit-token"})
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.location.endswith("/portal/fake-player-bearer"))
        limiter.assert_not_called()

    def test_reset_email_uses_untrusted_host_without_rate_limit(self):
        client = self.client()
        token = client.get("/login/csrf", headers={"Host": "attacker.example"}).json["token"]
        conn = MagicMock()
        conn.execute.return_value.fetchone.return_value = {"id": 1, "username": "admin", "email": "admin@example.invalid"}
        with patch.object(app, "get_connection", return_value=conn), patch.object(
            app, "crear_token_recuperacion", return_value="fake-reset-token"
        ), patch.object(app, "enviar_email", return_value=(True, None)) as email, patch.object(
            app, "registrar_auditoria"
        ), patch.object(app, "consumir_limite_publico") as limiter:
            response = client.post("/password/recuperar", headers={"Host": "attacker.example"},
                                   data={"email": "admin@example.invalid", "_csrf_token": token})
        self.assertEqual(response.status_code, 302)
        self.assertIn("http://attacker.example/password/restablecer/fake-reset-token", email.call_args.args[2])
        limiter.assert_not_called()

    def test_client_controls_ip_used_for_login_limits(self):
        with app.app.test_request_context("/login", headers={"X-Forwarded-For": "198.51.100.77, 203.0.113.2"},
                                          environ_base={"REMOTE_ADDR": "127.0.0.1"}):
            self.assertEqual(app.audit_request_ip(), "198.51.100.77")

    def test_anonymous_client_can_disable_another_push_endpoint(self):
        client = self.client()
        token = client.get("/login/csrf").json["token"]
        conn = MagicMock()
        with patch.object(app, "get_connection", return_value=conn), patch.object(
            app, "desactivar_suscripcion_push"
        ) as disable, patch.object(app, "registrar_auditoria"):
            response = client.post("/pwa/push/unsubscribe", json={"endpoint": "https://push.example.invalid/victim"},
                                   headers={"X-CSRF-Token": token})
        self.assertEqual(response.status_code, 200)
        disable.assert_called_once_with(conn, "https://push.example.invalid/victim")

    def test_push_repository_accepts_internal_http_endpoint(self):
        conn = MagicMock()
        notificaciones.guardar_suscripcion_push(conn,
            {"endpoint": "http://127.0.0.1:8080/private", "keys": {}},
            {"tipo": "usuario", "usuario_id": 1})
        self.assertIn("http://127.0.0.1:8080/private", conn.execute.call_args.args[1])

    def test_export_keeps_attacker_formula_as_formula(self):
        workbook = Workbook()
        app.append_fila_reporte(workbook.active, ['=HYPERLINK("https://example.invalid","abrir")'])
        self.assertEqual(workbook.active["A1"].data_type, "f")

    def test_reports_permission_alone_exports_medical_observations(self):
        client = self.client()
        with client.session_transaction() as state:
            state.update(user_id=1, username="reports-only", rol="custom-reports", permisos=["reportes_ver"])
        conn = MagicMock()
        def execute(query, params=None):
            rows = MagicMock()
            rows.fetchall.return_value = []
            if "LEFT JOIN fichas_medicas" in query:
                rows.fetchall.return_value = [{
                    "apellido": "Auditoria", "nombre": "Simulado", "categoria": "Plantel",
                    "presentada": 1, "apto_fisico": 1, "fecha_vencimiento": None,
                    "contacto_emergencia": "Simulado", "telefono_emergencia": "",
                    "documento_nombre": "", "ocr_fecha": None,
                    "observaciones": "AUDIT-PRIVATE-MEDICAL-NOTE",
                }]
            return rows
        conn.execute.side_effect = execute
        with tempfile.TemporaryDirectory(prefix="sig-security-") as directory, patch.object(
            app, "BASE_DIR", Path(directory)
        ), patch.object(app, "get_connection", return_value=conn), patch.object(
            app, "obtener_config_mantenimiento", return_value={"activo": False}
        ), patch.object(app, "registrar_auditoria"):
            response = client.get("/exportar/datos")
            self.assertEqual(response.status_code, 200)
            workbook = load_workbook(io.BytesIO(response.data))
            self.assertEqual(workbook["Fichas médicas"]["K2"].value, "AUDIT-PRIVATE-MEDICAL-NOTE")
            response.close()

    def test_upload_trusts_extension_instead_of_content(self):
        upload = FileStorage(stream=io.BytesIO(b"not a PDF"), filename="fake.pdf")
        validated = app.validar_comprobante_upload(upload)
        self.assertEqual(validated[3], "application/pdf")

    def test_audit_form_does_not_redact_tokens(self):
        with app.app.test_request_context("/", method="POST", data={"_csrf_token": "fake-csrf", "portal_token": "fake-bearer", "password": "fake-password"}):
            result = app.sanitized_audit_form()
        self.assertEqual(result["password"], "[redactado]")
        self.assertEqual(result["portal_token"], "fake-bearer")

    def test_non_ascii_csrf_raises_instead_of_rejecting(self):
        with app.app.test_request_context("/login", method="POST", data={"_csrf_token": "ñ"}):
            session["_csrf_token"] = "ascii-token"
            with self.assertRaises(TypeError):
                app.csrf_valido()

    def test_sensitive_download_lacks_no_store_and_security_headers(self):
        client = self.client()
        conn = MagicMock()
        conn.execute.return_value.fetchone.return_value = {"comprobante_drive_file_id": "fake-file", "comprobante_mime_type": "application/pdf", "comprobante_nombre": "fake.pdf"}
        with patch.object(app, "get_connection", return_value=conn), patch.object(
            app, "portal_tiene_notificaciones_activas", return_value=True
        ), patch.object(app, "descargar_drive_file", return_value=io.BytesIO(b"%PDF-1.4\n")):
            response = client.get("/portal/fake-bearer/cuotas/123/comprobante/ver")
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("no-store", response.headers.get("Cache-Control", ""))
        for header in ("X-Content-Type-Options", "X-Frame-Options", "Content-Security-Policy", "Referrer-Policy"):
            self.assertNotIn(header, response.headers)

    def test_csrf_still_rejects_missing_token_on_sensitive_post(self):
        with patch.object(app, "registrar_auditoria"), patch.object(app, "get_connection") as conn:
            self.assertEqual(self.client().post("/usuarios/123/eliminar").status_code, 400)
            conn.assert_not_called()


if __name__ == "__main__":
    unittest.main(verbosity=2)
